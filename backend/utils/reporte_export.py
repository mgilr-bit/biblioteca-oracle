"""Exportadores de reportes a XLSX (openpyxl) y CSV.

Reciben un `ReporteResponse` ya armado por `ReporteService` y solo se
ocupan de la presentación: no consultan la BD ni conocen las secciones.

Seguridad: los textos vienen de datos editables por usuarios (títulos,
detalles de auditoría, nombres), así que se neutraliza la inyección de
fórmulas (OWASP "CSV Injection") en ambos formatos.
"""
import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from schemas.reporte import ReporteColumna, ReporteResponse, ReporteTabla

# Paleta alineada con el frontend (tokens.css: --accent, --surface-alt, --border).
_COLOR_ACENTO = "2C3E58"
_COLOR_PANEL = "F2F0EA"
_COLOR_BORDE = "D4CFC0"
_COLOR_TENUE = "6D6A61"

_FORMATOS_EXCEL = {
    "entero": "#,##0",
    "decimal": "#,##0.00",
    "moneda": '"Q"#,##0.00',
    "fecha": "dd/mm/yyyy",
    "fechahora": "dd/mm/yyyy hh:mm",
}

_PREFIJOS_FORMULA = ("=", "+", "-", "@", "\t", "\r")
_CARACTERES_HOJA_INVALIDOS = re.compile(r"[\[\]:*?/\\]")


def nombre_archivo(reporte: ReporteResponse, extension: str) -> str:
    return f"reporte_{reporte.seccion}_{reporte.generado_en.strftime('%Y%m%d_%H%M')}.{extension}"


# --- Formateo común ---------------------------------------------------------

def _a_datetime(valor: Any) -> Any:
    if isinstance(valor, datetime) and valor.tzinfo is not None:
        # openpyxl no admite datetimes con zona horaria.
        return valor.replace(tzinfo=None)
    return valor


def _texto(valor: Any, tipo: str) -> str:
    """Representación legible de un valor para CSV/etiquetas."""
    if valor is None or valor == "":
        return ""
    if tipo == "fecha" and isinstance(valor, (datetime, date)):
        return valor.strftime("%d/%m/%Y")
    if tipo == "fechahora" and isinstance(valor, datetime):
        return valor.strftime("%d/%m/%Y %H:%M")
    if tipo == "moneda":
        return f"Q{Decimal(str(valor)):,.2f}"
    if tipo == "decimal":
        return f"{Decimal(str(valor)):,.2f}"
    if tipo == "entero":
        return f"{int(valor):,}"
    return str(valor)


def _valor_csv(valor: Any, tipo: str) -> str:
    if valor is None:
        return ""
    if tipo in ("entero", "decimal", "moneda"):
        # Números sin separador de miles para que las hojas de cálculo los reconozcan.
        return str(int(valor)) if tipo == "entero" else f"{Decimal(str(valor)):.2f}"
    texto = _texto(valor, tipo)
    if texto.startswith(_PREFIJOS_FORMULA):
        texto = "'" + texto
    return texto


# --- CSV --------------------------------------------------------------------

def exportar_csv(reporte: ReporteResponse) -> bytes:
    """CSV UTF-8 con BOM (Excel respeta tildes y eñes) y cabeceras legibles.

    Si el reporte tiene varias tablas, cada una va precedida de su título y
    separada por una línea en blanco.
    """
    salida = io.StringIO()
    writer = csv.writer(salida)
    varias = len(reporte.tablas) > 1

    for indice, tabla in enumerate(reporte.tablas):
        if varias:
            if indice:
                writer.writerow([])
            writer.writerow([tabla.titulo])
        writer.writerow([c.titulo for c in tabla.columnas])
        for fila in tabla.filas:
            writer.writerow([_valor_csv(fila.get(c.clave), c.tipo) for c in tabla.columnas])
        if tabla.totales:
            writer.writerow(_fila_totales(tabla, lambda v, c: _valor_csv(v, c.tipo)))

    return ("﻿" + salida.getvalue()).encode("utf-8")


def _fila_totales(tabla: ReporteTabla, formatear) -> list:
    fila = []
    for indice, columna in enumerate(tabla.columnas):
        if columna.clave in tabla.totales:
            fila.append(formatear(tabla.totales[columna.clave], columna))
        else:
            fila.append("Total" if indice == 0 else "")
    return fila


# --- XLSX -------------------------------------------------------------------

def _escribir(celda, valor: Any, tipo: str) -> None:
    if isinstance(valor, Decimal):
        valor = float(valor)
    valor = _a_datetime(valor)
    if isinstance(valor, str):
        valor = ILLEGAL_CHARACTERS_RE.sub("", valor)
    celda.value = valor
    if isinstance(valor, str) and valor.startswith("="):
        # openpyxl interpreta "=..." como fórmula: se fuerza a texto literal.
        celda.data_type = "s"
    formato = _FORMATOS_EXCEL.get(tipo)
    if formato and valor is not None and not isinstance(valor, str):
        celda.number_format = formato


def _nombre_hoja(titulo: str, usados: set[str]) -> str:
    base = " ".join(_CARACTERES_HOJA_INVALIDOS.sub(" ", titulo).split())[:31].strip() or "Datos"
    nombre, sufijo = base, 2
    while nombre.lower() in usados:
        extra = f" ({sufijo})"
        nombre = base[: 31 - len(extra)] + extra
        sufijo += 1
    usados.add(nombre.lower())
    return nombre


def _configurar_impresion(hoja, horizontal: bool) -> None:
    hoja.page_setup.orientation = "landscape" if horizontal else "portrait"
    hoja.page_setup.paperSize = hoja.PAPERSIZE_LETTER
    hoja.page_setup.fitToWidth = 1
    hoja.page_setup.fitToHeight = 0
    hoja.sheet_properties.pageSetUpPr.fitToPage = True
    hoja.oddFooter.center.text = "Página &P de &N"


def _hoja_resumen(hoja, reporte: ReporteResponse) -> None:
    negrita = Font(bold=True)
    tenue = Font(color=_COLOR_TENUE, size=9)
    seccion = Font(bold=True, color=_COLOR_ACENTO, size=12)
    borde_inferior = Border(bottom=Side(style="thin", color=_COLOR_BORDE))

    hoja["A1"] = "BIBLIOTECA · REPORTE"
    hoja["A1"].font = tenue
    hoja["A2"] = reporte.titulo
    hoja["A2"].font = Font(bold=True, size=18, color=_COLOR_ACENTO)
    fila = 3
    if reporte.subtitulo:
        hoja.cell(row=fila, column=1, value=reporte.subtitulo).font = Font(italic=True, color=_COLOR_TENUE)
        fila += 1

    def titulo_seccion(texto: str) -> None:
        nonlocal fila
        fila += 1
        for columna in (1, 2):
            hoja.cell(row=fila, column=columna).border = borde_inferior
        hoja.cell(row=fila, column=1, value=texto).font = seccion
        fila += 1

    titulo_seccion("Datos de emisión")
    for etiqueta, valor, tipo in (
        ("Generado el", reporte.generado_en, "fechahora"),
        ("Generado por", reporte.generado_por, "texto"),
    ):
        hoja.cell(row=fila, column=1, value=etiqueta).font = negrita
        _escribir(hoja.cell(row=fila, column=2), valor, tipo)
        hoja.cell(row=fila, column=2).alignment = Alignment(horizontal="left")
        fila += 1

    titulo_seccion("Filtros aplicados")
    if reporte.filtros:
        for filtro in reporte.filtros:
            hoja.cell(row=fila, column=1, value=filtro.etiqueta).font = negrita
            _escribir(hoja.cell(row=fila, column=2), filtro.valor, "texto")
            fila += 1
    else:
        hoja.cell(row=fila, column=1, value="Sin filtros adicionales").font = tenue
        fila += 1

    if reporte.indicadores:
        titulo_seccion("Indicadores")
        for indicador in reporte.indicadores:
            hoja.cell(row=fila, column=1, value=indicador.etiqueta).font = negrita
            celda = hoja.cell(row=fila, column=2)
            _escribir(celda, indicador.valor, indicador.tipo)
            celda.alignment = Alignment(horizontal="left")
            fila += 1

    if reporte.tablas:
        titulo_seccion("Contenido")
        for tabla in reporte.tablas:
            hoja.cell(row=fila, column=1, value=tabla.titulo)
            hoja.cell(row=fila, column=2, value=f"{len(tabla.filas):,} registro(s)")
            fila += 1

    if reporte.notas:
        titulo_seccion("Notas")
        for nota in reporte.notas:
            hoja.cell(row=fila, column=1, value=nota).font = tenue
            fila += 1

    hoja.column_dimensions["A"].width = 34
    hoja.column_dimensions["B"].width = 48
    hoja.sheet_view.showGridLines = False
    _configurar_impresion(hoja, horizontal=False)


def _ancho(columna: ReporteColumna, tabla: ReporteTabla) -> float:
    muestras = [len(columna.titulo)]
    for fila in tabla.filas[:500]:
        muestras.append(len(_texto(fila.get(columna.clave), columna.tipo)))
    return min(max(max(muestras) + 2, 9), 60)


def _hoja_tabla(hoja, reporte: ReporteResponse, tabla: ReporteTabla) -> None:
    columnas = tabla.columnas
    ultima = get_column_letter(len(columnas))

    hoja["A1"] = tabla.titulo
    hoja["A1"].font = Font(bold=True, size=14, color=_COLOR_ACENTO)
    hoja["A2"] = (
        f"{reporte.titulo} · Generado el {reporte.generado_en.strftime('%d/%m/%Y %H:%M')} "
        f"por {reporte.generado_por}"
    )
    hoja["A2"].font = Font(color=_COLOR_TENUE, size=9)

    fila_cabecera = 4
    relleno_cabecera = PatternFill("solid", fgColor=_COLOR_ACENTO)
    for indice, columna in enumerate(columnas, start=1):
        celda = hoja.cell(row=fila_cabecera, column=indice, value=columna.titulo)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = relleno_cabecera
        celda.alignment = Alignment(
            vertical="center",
            horizontal="right" if columna.tipo in ("entero", "decimal", "moneda") else "left",
        )
        hoja.column_dimensions[get_column_letter(indice)].width = _ancho(columna, tabla)
    hoja.row_dimensions[fila_cabecera].height = 20

    fila = fila_cabecera
    for registro in tabla.filas:
        fila += 1
        for indice, columna in enumerate(columnas, start=1):
            _escribir(hoja.cell(row=fila, column=indice), registro.get(columna.clave), columna.tipo)

    if not tabla.filas:
        fila += 1
        hoja.cell(row=fila, column=1, value="Sin registros para los filtros aplicados").font = Font(
            italic=True, color=_COLOR_TENUE
        )
    else:
        hoja.auto_filter.ref = f"A{fila_cabecera}:{ultima}{fila}"

    if tabla.totales:
        fila += 1
        relleno_total = PatternFill("solid", fgColor=_COLOR_PANEL)
        borde_total = Border(top=Side(style="thin", color=_COLOR_ACENTO))
        for indice, valor in enumerate(_fila_totales(tabla, lambda v, c: v), start=1):
            celda = hoja.cell(row=fila, column=indice)
            _escribir(celda, valor if valor != "" else None, columnas[indice - 1].tipo)
            celda.font = Font(bold=True)
            celda.fill = relleno_total
            celda.border = borde_total

    hoja.freeze_panes = f"A{fila_cabecera + 1}"
    hoja.print_title_rows = f"{fila_cabecera}:{fila_cabecera}"
    _configurar_impresion(hoja, horizontal=len(columnas) > 5)


def exportar_xlsx(reporte: ReporteResponse) -> bytes:
    """Libro con una hoja "Resumen" (emisión, filtros, indicadores) y una hoja
    por tabla con cabecera fija, autofiltro, formatos numéricos y totales."""
    libro = Workbook()
    libro.properties.title = reporte.titulo
    libro.properties.creator = reporte.generado_por

    resumen = libro.active
    resumen.title = "Resumen"
    _hoja_resumen(resumen, reporte)

    usados = {"resumen"}
    for tabla in reporte.tablas:
        _hoja_tabla(libro.create_sheet(_nombre_hoja(tabla.titulo, usados)), reporte, tabla)

    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()
